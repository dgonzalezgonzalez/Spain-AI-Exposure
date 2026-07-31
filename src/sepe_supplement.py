from __future__ import annotations

from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook


MONTHS = {
    "ENERO": "01", "FEBRERO": "02", "MARZO": "03", "ABRIL": "04",
    "MAYO": "05", "JUNIO": "06", "JULIO": "07", "AGOSTO": "08",
    "SEPTIEMBRE": "09", "OCTUBRE": "10", "NOVIEMBRE": "11", "DICIEMBRE": "12",
}
PROVINCES = {
    1: "Araba/\u00c1lava", 2: "Albacete", 3: "Alicante/Alacant", 4: "Almer\u00eda", 5: "\u00c1vila",
    6: "Badajoz", 7: "Balears, Illes", 8: "Barcelona", 9: "Burgos", 10: "C\u00e1ceres",
    11: "C\u00e1diz", 12: "Castell\u00f3n/Castell\u00f3", 13: "Ciudad Real", 14: "C\u00f3rdoba", 15: "Coru\u00f1a, A",
    16: "Cuenca", 17: "Girona", 18: "Granada", 19: "Guadalajara", 20: "Gipuzkoa",
    21: "Huelva", 22: "Huesca", 23: "Ja\u00e9n", 24: "Le\u00f3n", 25: "Lleida", 26: "Rioja, La",
    27: "Lugo", 28: "Madrid", 29: "M\u00e1laga", 30: "Murcia", 31: "Navarra", 32: "Ourense",
    33: "Asturias", 34: "Palencia", 35: "Palmas, Las", 36: "Pontevedra", 37: "Salamanca",
    38: "Santa Cruz de Tenerife", 39: "Cantabria", 40: "Segovia", 41: "Sevilla", 42: "Soria",
    43: "Tarragona", 44: "Teruel", 45: "Toledo", 46: "Valencia/Val\u00e8ncia", 47: "Valladolid",
    48: "Bizkaia", 49: "Zamora", 50: "Zaragoza", 51: "Ceuta", 52: "Melilla",
}
AGE_GROUPS = ["<18", "18-24", "25-29", "30-39", "40-44", ">44"]
REQUIRED_COLUMNS = ["period", "cno4", "dimension", "category", "gender", "contratos", "parados", "personas", "source_url"]


def _number(value: object) -> int:
    if value is None:
        return 0
    if isinstance(value, str):
        cleaned = value.strip().lstrip("'").strip()
        if cleaned in {"", "-"}:
            return 0
        cleaned = cleaned.replace(".", "").replace(",", ".")
        return int(float(cleaned))
    return int(value)


def parse_provider_workbook(
    workbook_path: Path,
    report_keys: Iterable[tuple[str, str]],
) -> pd.DataFrame:
    """Extract complete gender, age, and province breakdowns for selected reports."""
    requested = {(str(cno4).zfill(4), period) for cno4, period in report_keys}
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    output: list[dict[str, object]] = []
    try:
        sheets = {name[:4]: name for name in workbook.sheetnames if name[:4].isdigit()}
        for cno4 in sorted({key[0] for key in requested}):
            sheet = workbook[sheets[cno4]]
            headers = list(sheet.iter_rows(min_row=4, max_row=6, values_only=True))
            columns: list[tuple[int, str, str, str]] = []
            current_month = None
            current_gender = None
            for index in range(3, sheet.max_column):
                if headers[0][index]:
                    current_month = str(headers[0][index]).strip().upper()
                if headers[1][index]:
                    current_gender = str(headers[1][index]).strip().title()
                measure_label = str(headers[2][index] or "").strip().lower()
                measure = "contratos" if "contratos" in measure_label else "parados" if "parados" in measure_label else ""
                if current_month and current_gender and measure:
                    month_word = current_month.split()[0]
                    year = current_month.split()[-1]
                    period = f"{year}-{MONTHS[month_word]}"
                    if (cno4, period) in requested:
                        columns.append((index, period, current_gender, measure))

            atoms: list[dict[str, object]] = []
            province_id = None
            for row in sheet.iter_rows(min_row=7, values_only=True):
                if row[0] is not None:
                    province_id = int(row[0])
                age = str(row[2] or "").strip()
                if age not in AGE_GROUPS or province_id not in PROVINCES:
                    continue
                for index, period, gender, measure in columns:
                    atoms.append({
                        "period": period, "province_id": province_id, "age": age,
                        "gender": gender, "measure": measure, "value": _number(row[index]),
                    })
            atomic = pd.DataFrame(atoms)
            for period in sorted(atomic["period"].unique()):
                subset = atomic[atomic["period"] == period]
                source = f"provider://sepe/ocup_prov_2022.xlsx#{cno4}/{period}"
                frames = [
                    ("gender", "gender", ["Hombre", "Mujer"]),
                    ("age", "age", AGE_GROUPS),
                    ("province", "province_id", list(PROVINCES)),
                ]
                for dimension, grouping, categories in frames:
                    grouped = subset.groupby([grouping, "measure"])["value"].sum().unstack(fill_value=0)
                    grouped = grouped.reindex(categories, fill_value=0)
                    for category, values in grouped.iterrows():
                        label = PROVINCES[category] if dimension == "province" else category
                        output.append({
                            "period": period, "cno4": cno4, "dimension": dimension,
                            "category": label, "gender": label if dimension == "gender" else "Total",
                            "contratos": int(values.get("contratos", 0)),
                            "parados": int(values.get("parados", 0)), "personas": pd.NA,
                            "source_url": source,
                        })
    finally:
        workbook.close()
    result = pd.DataFrame(output, columns=REQUIRED_COLUMNS)
    validate_supplement(result)
    return result


def validate_supplement(supplement: pd.DataFrame) -> None:
    missing = set(REQUIRED_COLUMNS) - set(supplement.columns)
    if missing:
        raise ValueError(f"SEPE supplement is missing columns: {sorted(missing)}")
    expected = {"gender": 2, "age": 6, "province": 52}
    for key, report in supplement.groupby(["period", "cno4"]):
        counts = report.groupby("dimension").size().to_dict()
        if counts != expected:
            raise ValueError(f"Incomplete SEPE supplement for {key}: {counts}")
        totals = report.groupby("dimension")[["contratos", "parados"]].sum()
        if not totals.eq(totals.iloc[0]).all().all():
            raise ValueError(f"Breakdown totals disagree for {key}: {totals.to_dict()}")


def load_supplement(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame(columns=REQUIRED_COLUMNS)
    supplement = pd.read_csv(path, dtype={"cno4": str})
    supplement["cno4"] = supplement["cno4"].str.zfill(4)
    validate_supplement(supplement)
    return supplement


def add_missing_breakdowns(report: pd.DataFrame, supplement: pd.DataFrame) -> pd.DataFrame:
    if report.empty or supplement.empty:
        return report
    cno4 = str(report.iloc[0]["cno4"]).zfill(4)
    period = str(report.iloc[0]["period"])
    extra = supplement[(supplement["cno4"] == cno4) & (supplement["period"] == period)].copy()
    if extra.empty:
        return report
    total = report[(report["dimension"] == "total") & (report["category"] == "Total")]
    if len(total) != 1:
        raise ValueError(f"Expected one total row for {(cno4, period)}")
    for measure in ("contratos", "parados"):
        expected = total.iloc[0][measure]
        actual = extra[extra["dimension"] == "gender"][measure].sum()
        if pd.notna(expected) and int(expected) != int(actual):
            raise ValueError(f"Provider {measure} total {actual} disagrees with report total {expected} for {(cno4, period)}")
    present_dimensions = set(report["dimension"])
    extra = extra[~extra["dimension"].isin(present_dimensions)]
    if extra.empty:
        return report
    extra["occupation_title"] = report.iloc[0]["occupation_title"]
    return pd.concat([report, extra[report.columns]], ignore_index=True)
